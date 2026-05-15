"""
import_word.py — парсинг Word-файлов и импорт в PostgreSQL
Использует Claude API для извлечения структурированных данных из «киселя»

Установка:
  pip install python-docx anthropic psycopg2-binary python-dotenv

Запуск:
  python import_word.py --file data/commanders.docx --type persons
  python import_word.py --file data/combat_path.docx --type events
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

import anthropic
import psycopg2
from docx import Document
from dotenv import load_dotenv

load_dotenv()

# ── Подключение к БД ──────────────────────────────────────────────────────
def get_db():
    return psycopg2.connect(
        dbname=os.getenv('POSTGRES_DB',       'diviziya52'),
        user=os.getenv('POSTGRES_USER',        'diviziya52_user'),
        password=os.getenv('POSTGRES_PASSWORD',''),
        host=os.getenv('POSTGRES_HOST',        'localhost'),
        port=os.getenv('POSTGRES_PORT',        '5432'),
    )

# ── Извлечение текста из Word ─────────────────────────────────────────────
def extract_text(path: str) -> str:
    doc = Document(path)
    chunks = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # сохраняем стиль как подсказку для Claude
            style = para.style.name
            if 'Heading' in style:
                chunks.append(f"\n=== {text} ===\n")
            else:
                chunks.append(text)
    # таблицы тоже
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)

# ── Claude API — извлечение персон ───────────────────────────────────────
PERSONS_PROMPT = """Ты помогаешь оцифровать архив 52-й стрелковой дивизии СССР.

Из текста ниже извлеки всех упомянутых людей (офицеров и солдат).
Верни ТОЛЬКО валидный JSON — массив объектов, без пояснений, без markdown.

Структура каждого объекта:
{
  "last_name": "обязательно",
  "first_name": "или null",
  "patronymic": "или null",
  "rank": "воинское звание или null",
  "role": "должность/роль или null",
  "unit": "подразделение или null",
  "date_from": "YYYY-MM-DD или YYYY-MM или YYYY или null",
  "date_to":   "YYYY-MM-DD или YYYY-MM или YYYY или null",
  "awards": ["награда1", "награда2"],
  "birth_place": "или null",
  "note": "любая дополнительная информация или null"
}

Текст:
"""

EVENTS_PROMPT = """Ты помогаешь оцифровать боевой путь 52-й стрелковой дивизии СССР.

Из текста ниже извлеки все события (перемещения, бои, операции).
Верни ТОЛЬКО валидный JSON — массив объектов, без пояснений, без markdown.

Структура каждого объекта:
{
  "title": "краткое название события, обязательно",
  "date_from": "YYYY-MM-DD или YYYY-MM или YYYY, обязательно",
  "date_to":   "YYYY-MM-DD или null если одна дата",
  "location":  "название места или null",
  "event_type": "battle|move|defense|offense|reorg|command|other",
  "is_key": true/false,
  "description": "подробное описание или null",
  "commanders": ["Фамилия И.О.", ...],
  "casualties_note": "упоминания о потерях или null"
}

Текст:
"""

def extract_with_claude(text: str, extract_type: str) -> list:
    client = anthropic.Anthropic()  # берёт ANTHROPIC_API_KEY из окружения

    prompt = PERSONS_PROMPT if extract_type == 'persons' else EVENTS_PROMPT

    # Разбиваем большой текст на куски по ~3000 символов с перекрытием
    chunk_size = 3000
    overlap    = 300
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunks.append(text[start:end])
        start += chunk_size - overlap

    all_results = []
    for i, chunk in enumerate(chunks):
        print(f"  Обрабатываю кусок {i+1}/{len(chunks)}...", end=' ')
        try:
            response = client.messages.create(
                model="claude-sonnet-4-20250514",
                max_tokens=2000,
                messages=[{"role": "user", "content": prompt + chunk}]
            )
            raw = response.content[0].text.strip()
            # убираем возможные markdown-обёртки
            raw = re.sub(r'^```json\s*', '', raw)
            raw = re.sub(r'\s*```$',     '', raw)
            data = json.loads(raw)
            if isinstance(data, list):
                all_results.extend(data)
                print(f"найдено {len(data)} записей")
            else:
                print("неожиданный формат, пропускаю")
        except (json.JSONDecodeError, Exception) as e:
            print(f"ошибка: {e}")
            continue

    return all_results

# ── Импорт персон ─────────────────────────────────────────────────────────
def import_persons(records: list, db):
    cur = db.cursor()
    inserted = 0
    skipped  = 0

    for r in records:
        if not r.get('last_name'):
            skipped += 1
            continue

        # Проверка дубликатов по фамилии + имени
        cur.execute("""
            SELECT id FROM persons
            WHERE last_name = %s AND (first_name = %s OR first_name IS NULL)
            LIMIT 1
        """, (r['last_name'], r.get('first_name')))
        existing = cur.fetchone()

        if existing:
            skipped += 1
            continue

        # Вставка персоны
        cur.execute("""
            INSERT INTO persons (last_name, first_name, patronymic, awards)
            VALUES (%s, %s, %s, %s)
            RETURNING id
        """, (
            r['last_name'],
            r.get('first_name'),
            r.get('patronymic'),
            r.get('awards') or [],
        ))
        person_id = cur.fetchone()[0]

        # Запись о службе если есть данные о должности
        if r.get('rank') or r.get('role'):
            cur.execute("""
                INSERT INTO service_records
                  (person_id, rank, role, date_from, date_to, note)
                VALUES (%s, %s, %s, %s, %s, %s)
            """, (
                person_id,
                r.get('rank'),
                r.get('role'),
                _parse_date(r.get('date_from')),
                _parse_date(r.get('date_to')),
                r.get('note'),
            ))

        inserted += 1

    db.commit()
    cur.close()
    return inserted, skipped

# ── Импорт событий ────────────────────────────────────────────────────────
def import_events(records: list, db, division_id: int = 1):
    cur = db.cursor()
    inserted = 0
    skipped  = 0

    for r in records:
        if not r.get('title') or not r.get('date_from'):
            skipped += 1
            continue

        cur.execute("""
            INSERT INTO events
              (division_id, date_from, date_to, title, description,
               location, event_type, is_key)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """, (
            division_id,
            _parse_date(r['date_from']),
            _parse_date(r.get('date_to')),
            r['title'],
            r.get('description'),
            r.get('location'),
            r.get('event_type', 'other'),
            bool(r.get('is_key', False)),
        ))
        inserted += 1

    db.commit()
    cur.close()
    return inserted, skipped

# ── Вспомогательные функции ───────────────────────────────────────────────
def _parse_date(s):
    """Преобразует строку 'YYYY', 'YYYY-MM', 'YYYY-MM-DD' в date или None."""
    if not s:
        return None
    s = str(s).strip()
    if re.match(r'^\d{4}$', s):
        return f"{s}-01-01"
    if re.match(r'^\d{4}-\d{2}$', s):
        return f"{s}-01"
    if re.match(r'^\d{4}-\d{2}-\d{2}$', s):
        return s
    return None

# ── CLI ───────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description='Импорт данных из Word в PostgreSQL')
    parser.add_argument('--file', required=True, help='Путь к .docx файлу')
    parser.add_argument('--type', required=True, choices=['persons', 'events'],
                        help='Тип данных: persons или events')
    parser.add_argument('--dry-run', action='store_true',
                        help='Только показать что будет импортировано, не записывать')
    args = parser.parse_args()

    if not Path(args.file).exists():
        print(f"Файл не найден: {args.file}")
        sys.exit(1)

    print(f"Читаю {args.file}...")
    text = extract_text(args.file)
    print(f"Извлечено {len(text)} символов текста")

    print(f"Отправляю в Claude для структурирования ({args.type})...")
    records = extract_with_claude(text, args.type)
    print(f"Всего записей от Claude: {len(records)}")

    if args.dry_run:
        print("\n--- DRY RUN (первые 5 записей) ---")
        for r in records[:5]:
            print(json.dumps(r, ensure_ascii=False, indent=2))
        return

    print("Подключаюсь к БД...")
    db = get_db()

    if args.type == 'persons':
        ins, skp = import_persons(records, db)
    else:
        ins, skp = import_events(records, db)

    db.close()
    print(f"\nГотово: вставлено {ins}, пропущено {skp}")

if __name__ == '__main__':
    main()
